"""
doc_builder.py  — Helper utilities for the report generator.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_page_margins(doc, top=2, bottom=2, left=2.5, right=2.5):
    section = doc.sections[0]
    section.top_margin    = Cm(top)
    section.bottom_margin = Cm(bottom)
    section.left_margin   = Cm(left)
    section.right_margin  = Cm(right)


def add_page_number(doc):
    """Add centred page numbers to footer."""
    section = doc.sections[0]
    footer  = section.footer
    para    = footer.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = para.add_run()
    fld = OxmlElement("w:fldChar")
    fld.set(qn("w:fldCharType"), "begin")
    run._r.append(fld)
    run2 = para.add_run()
    ins  = OxmlElement("w:instrText")
    ins.text = "PAGE"
    run2._r.append(ins)
    run3 = para.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "end")
    run3._r.append(fld2)


def cover(doc, title, subtitle, author, date):
    doc.add_paragraph()
    doc.add_paragraph()
    doc.add_paragraph()
    t = doc.add_paragraph(title)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    t.runs[0].font.size = Pt(28)
    t.runs[0].bold = True
    t.runs[0].font.color.rgb = RGBColor(0x1F, 0x35, 0x64)

    s = doc.add_paragraph(subtitle)
    s.alignment = WD_ALIGN_PARAGRAPH.CENTER
    s.runs[0].font.size = Pt(14)
    s.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)
    s.runs[0].italic = True

    doc.add_paragraph()
    a = doc.add_paragraph(f"Author: {author}   |   Date: {date}")
    a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    a.runs[0].font.size = Pt(11)
    doc.add_page_break()


def h1(doc, text):
    p = doc.add_heading("", level=1)
    run = p.add_run(text)
    run.font.size = Pt(16)
    run.font.color.rgb = RGBColor(0x1F, 0x35, 0x64)
    run.bold = True
    return p


def h2(doc, text):
    p = doc.add_heading("", level=2)
    run = p.add_run(text)
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x2E, 0x74, 0xB5)
    run.bold = True
    return p


def h3(doc, text):
    p = doc.add_heading("", level=3)
    run = p.add_run(text)
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(0x20, 0x60, 0x20)
    run.bold = True
    return p


def body(doc, text):
    p = doc.add_paragraph(text)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    for run in p.runs:
        run.font.size = Pt(11)
    return p


def code(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent  = Inches(0.3)
    p.paragraph_format.right_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    shading = OxmlElement("w:shd")
    shading.set(qn("w:val"),   "clear")
    shading.set(qn("w:color"), "auto")
    shading.set(qn("w:fill"),  "EEF0F4")
    p._p.pPr.append(shading) if p._p.pPr is not None else None
    run = p.add_run(text)
    run.font.name = "Courier New"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x5E)
    return p


def note(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after  = Pt(0)
    run = p.add_run("💡 " + text)
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = RGBColor(0x5A, 0x5A, 0x5A)
    return p


def divider(doc):
    doc.add_paragraph("─" * 72)


def two_col_table(doc, rows, bold_header=True):
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(rows):
        t.rows[i].cells[0].text = k
        t.rows[i].cells[1].text = v
        if i == 0 and bold_header:
            for cell in t.rows[i].cells:
                for run in cell.paragraphs[0].runs:
                    run.bold = True
    doc.add_paragraph()
    return t
